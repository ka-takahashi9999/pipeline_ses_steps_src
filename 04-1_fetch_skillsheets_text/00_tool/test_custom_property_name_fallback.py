#!/usr/bin/env python3
"""custom propertyのname欠損だけを救済するXLSX fallbackのFocused Test。"""

import base64
import hashlib
import importlib.util
import json
import sys
import unittest
import zipfile
import xml.etree.ElementTree as ET
from io import BytesIO
from pathlib import Path
from unittest.mock import patch


TOOL_DIR = Path(__file__).resolve().parent
STEP_DIR = TOOL_DIR.parent
ROOT = STEP_DIR.parent
sys.path.insert(0, str(ROOT))


def _load_module():
    path = TOOL_DIR / "fetch_skillsheets_text.py"
    spec = importlib.util.spec_from_file_location("fetch_custom_property_test", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


fetch = _load_module()
TARGET_ID = "s3_fb2f195a96a2be54f8e9611683d48a62158645259675c15e812752cd4a92d7f2"
NORMAL_ID = "s3_7173d3a335b7a70eebce8ab6ea3f8e957ceb5125cb0e97a28adad52a28f6b4f5"


def _read_mail(message_id):
    with fetch.INPUT_MAIL_MASTER.open(encoding="utf-8") as source:
        for line in source:
            record = json.loads(line)
            if record.get("message_id") == message_id:
                return record
    raise AssertionError(f"mail not found: {message_id}")


def _attachment_bytes(attachment):
    encoded = attachment["data"]
    return base64.urlsafe_b64decode(encoded + "=" * (-len(encoded) % 4))


def _xlsx_attachment(mail):
    return next(
        attachment
        for attachment in mail.get("attachments") or []
        if (attachment.get("filename") or "").lower().endswith(".xlsx")
    )


def _replace_zip_member(data, member_name, replacement):
    output = BytesIO()
    with zipfile.ZipFile(BytesIO(data), "r") as source:
        with zipfile.ZipFile(output, "w") as destination:
            destination.comment = source.comment
            for entry in source.infolist():
                value = replacement if entry.filename == member_name else source.read(entry)
                destination.writestr(entry, value)
    return output.getvalue()


class CustomPropertyNameFallbackTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.target_mail = _read_mail(TARGET_ID)
        cls.target_attachment = _xlsx_attachment(cls.target_mail)
        cls.target_data = _attachment_bytes(cls.target_attachment)
        cls.normal_mail = _read_mail(NORMAL_ID)
        cls.normal_attachment = _xlsx_attachment(cls.normal_mail)
        cls.normal_data = _attachment_bytes(cls.normal_attachment)

    def test_target_retries_once_and_passes_existing_quality(self):
        import openpyxl

        original_sha256 = hashlib.sha256(self.target_data).hexdigest()
        with patch.object(
            openpyxl, "load_workbook", wraps=openpyxl.load_workbook
        ) as load_workbook:
            text = fetch.extract_text_from_excel(
                self.target_data, self.target_attachment["filename"]
            )

        self.assertEqual(load_workbook.call_count, 2)
        self.assertEqual(text.count("=== シート:"), 2)
        self.assertEqual(len(text), 13110)
        self.assertIsNone(fetch.validate_skillsheet_text(text))
        self.assertEqual(hashlib.sha256(self.target_data).hexdigest(), original_sha256)

        result = fetch.fetch_skillsheet(
            TARGET_ID,
            self.target_mail,
            {"body_text": self.target_mail.get("body_text", "")},
        )
        self.assertTrue(result["success"])
        self.assertEqual(result["source"], "attachment")
        self.assertEqual(result["skillsheet"], text)

        sanitized = fetch._sanitize_missing_custom_property_names(self.target_data)
        self.assertIsNotNone(sanitized)
        workbook = openpyxl.load_workbook(
            BytesIO(sanitized), data_only=True, read_only=True
        )
        nonempty_cells = sum(
            1
            for worksheet in workbook.worksheets
            for row in worksheet.iter_rows(values_only=True)
            for value in row
            if value is not None and str(value).strip()
        )
        self.assertEqual((len(workbook.sheetnames), nonempty_cells), (2, 953))

    def test_existing_normal_xlsx_uses_one_load_and_text_is_unchanged(self):
        import openpyxl

        with fetch.OUTPUT_JSONL.open(encoding="utf-8") as source:
            current = next(
                record
                for record in (json.loads(line) for line in source)
                if record.get("message_id") == NORMAL_ID
            )
        with patch.object(
            fetch, "_sanitize_missing_custom_property_names", wraps=fetch._sanitize_missing_custom_property_names
        ) as sanitize, patch.object(
            openpyxl, "load_workbook", wraps=openpyxl.load_workbook
        ) as load_workbook:
            text = fetch.extract_text_from_excel(
                self.normal_data, self.normal_attachment["filename"]
            )

        self.assertEqual(load_workbook.call_count, 1)
        sanitize.assert_not_called()
        self.assertEqual(text, current["skillsheet"])

    def test_mixed_properties_remove_only_missing_name(self):
        with zipfile.ZipFile(BytesIO(self.target_data), "r") as source:
            custom_root = ET.fromstring(source.read(fetch.XLSX_CUSTOM_PROPERTIES))

        valid_property = ET.SubElement(
            custom_root,
            f"{{{fetch.CUSTOM_PROPERTIES_NS}}}property",
            {
                "fmtid": "{D5CDD505-2E9C-101B-9397-08002B2CF9AE}",
                "pid": "256",
                "name": "KeepMe",
            },
        )
        value = ET.SubElement(
            valid_property,
            f"{{{fetch.CUSTOM_PROPERTIES_VALUE_NS}}}lpwstr",
        )
        value.text = "preserved"
        mixed = _replace_zip_member(
            self.target_data,
            fetch.XLSX_CUSTOM_PROPERTIES,
            ET.tostring(custom_root, encoding="utf-8", xml_declaration=True),
        )

        sanitized = fetch._sanitize_missing_custom_property_names(mixed)
        self.assertIsNotNone(sanitized)
        with zipfile.ZipFile(BytesIO(mixed), "r") as before, zipfile.ZipFile(
            BytesIO(sanitized), "r"
        ) as after:
            self.assertEqual(before.namelist(), after.namelist())
            for name in before.namelist():
                if name != fetch.XLSX_CUSTOM_PROPERTIES:
                    self.assertEqual(before.read(name), after.read(name), name)
            root = ET.fromstring(after.read(fetch.XLSX_CUSTOM_PROPERTIES))

        properties = [
            child
            for child in list(root)
            if child.tag == f"{{{fetch.CUSTOM_PROPERTIES_NS}}}property"
        ]
        self.assertEqual([prop.get("name") for prop in properties], ["KeepMe"])
        self.assertEqual(
            properties[0].find(
                f"{{{fetch.CUSTOM_PROPERTIES_VALUE_NS}}}lpwstr"
            ).text,
            "preserved",
        )

    def test_other_corruption_is_not_repaired(self):
        cases_without_custom_inspection = {
            "broken_zip": b"PK\x03\x04broken",
            "broken_workbook": _replace_zip_member(
                self.target_data, "xl/workbook.xml", b"<workbook>"
            ),
            "other_custom_xml_error": _replace_zip_member(
                self.target_data, fetch.XLSX_CUSTOM_PROPERTIES, b"<Properties>"
            ),
        }
        for label, data in cases_without_custom_inspection.items():
            with self.subTest(label=label), patch.object(
                fetch,
                "_sanitize_missing_custom_property_names",
                wraps=fetch._sanitize_missing_custom_property_names,
            ) as sanitize:
                with self.assertRaises(Exception):
                    fetch.extract_text_from_excel(data, "fixture.xlsx")
                sanitize.assert_not_called()

        # read_only loadはcustom propertyを先に読むため、worksheet破損の場合は
        # 限定条件の構造検査まで進むが、sanitized copyの再読込は行わない。
        broken_worksheet = _replace_zip_member(
            self.target_data, "xl/worksheets/sheet1.xml", b"<worksheet>"
        )
        import openpyxl

        with patch.object(
            openpyxl, "load_workbook", wraps=openpyxl.load_workbook
        ) as load_workbook, patch.object(
            fetch,
            "_sanitize_missing_custom_property_names",
            wraps=fetch._sanitize_missing_custom_property_names,
        ) as sanitize:
            with self.assertRaises(Exception):
                fetch.extract_text_from_excel(broken_worksheet, "fixture.xlsx")
            sanitize.assert_called_once()
            self.assertEqual(load_workbook.call_count, 1)

    def test_sanitizer_rejects_required_structure_damage(self):
        damaged = _replace_zip_member(
            self.target_data, "xl/worksheets/sheet1.xml", b"<worksheet>"
        )
        self.assertIsNone(fetch._sanitize_missing_custom_property_names(damaged))
        wrong_workbook_root = _replace_zip_member(
            self.target_data, "xl/workbook.xml", b"<not-a-workbook/>"
        )
        self.assertIsNone(
            fetch._sanitize_missing_custom_property_names(wrong_workbook_root)
        )
        self.assertIsNone(
            fetch._sanitize_missing_custom_property_names(b"PK\x03\x04broken")
        )


if __name__ == "__main__":
    unittest.main()
